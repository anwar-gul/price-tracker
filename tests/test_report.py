from __future__ import annotations

import csv

import pytest
from openpyxl import load_workbook

from pricetracker import db
from pricetracker.report import export_all, print_report


@pytest.fixture
def seeded_conn(tmp_path):
    conn = db.connect(tmp_path / "seed.db")
    db.init_schema(conn)
    db.upsert_book(
        conn, url="https://example.com/a", title="Alpha", price=10.0,
        rating=5, in_stock=True, scraped_at="2026-03-01T00:00:00+00:00",
    )
    db.upsert_book(
        conn, url="https://example.com/b", title="Beta", price=20.0,
        rating=3, in_stock=False, scraped_at="2026-03-01T00:00:00+00:00",
    )
    db.upsert_book(
        conn, url="https://example.com/a", title="Alpha", price=8.0,
        rating=5, in_stock=True, scraped_at="2026-03-02T00:00:00+00:00",
    )
    yield conn
    conn.close()


def test_export_creates_csvs_and_workbook(tmp_path, seeded_conn):
    paths = export_all(seeded_conn, tmp_path / "exports")
    for key in ("catalog_csv", "history_csv", "workbook"):
        assert paths[key]
    assert (tmp_path / "exports" / "catalog.csv").exists()
    assert (tmp_path / "exports" / "price_history.csv").exists()
    assert (tmp_path / "exports" / "catalog.xlsx").exists()


def test_catalog_csv_contents(tmp_path, seeded_conn):
    paths = export_all(seeded_conn, tmp_path / "exports")
    with open(paths["catalog_csv"], newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 2
    titles = {row["Title"] for row in rows}
    assert titles == {"Alpha", "Beta"}


def test_workbook_sheets_and_change_row(tmp_path, seeded_conn):
    paths = export_all(seeded_conn, tmp_path / "exports")
    wb = load_workbook(paths["workbook"])
    assert wb.sheetnames == ["Catalog", "Price Changes", "Price Trends"]

    ws = wb["Price Changes"]
    assert ws.max_row == 2  # header + 1 change
    row = [cell.value for cell in ws[2]]
    assert row[1] == "Alpha"
    assert row[2] == 10.0
    assert row[3] == 8.0
    assert row[4] == -2.0
    assert row[5] == -20.0


def test_trend_sheet_data_and_chart(tmp_path, seeded_conn):
    paths = export_all(seeded_conn, tmp_path / "exports")
    wb = load_workbook(paths["workbook"])
    ws = wb["Price Trends"]

    assert [cell.value for cell in ws[1]] == ["Date", "Alpha"]
    assert ws["B2"].value == 10.0  # Alpha day 1
    assert ws["B3"].value == 8.0  # Alpha day 2
    assert ws["B2"].number_format == "#,##0.00"

    assert len(ws._charts) == 1
    chart = ws._charts[0]
    assert chart.title is not None
    assert chart.display_blanks == "span"
    assert len(chart.series) == 1


def test_trend_sheet_skipped_without_changes(tmp_path):
    conn = db.connect(tmp_path / "flat.db")
    db.init_schema(conn)
    db.upsert_book(
        conn, url="https://example.com/a", title="Alpha", price=10.0,
        rating=5, in_stock=True, scraped_at="2026-03-01T00:00:00+00:00",
    )
    wb = load_workbook(export_all(conn, tmp_path / "exports")["workbook"])
    conn.close()
    assert wb.sheetnames == ["Catalog", "Price Changes"]


def test_trend_sheet_caps_product_columns(tmp_path):
    conn = db.connect(tmp_path / "many.db")
    db.init_schema(conn)
    for i in range(14):
        url = f"https://example.com/{i}"
        db.upsert_book(conn, url=url, title=f"Item {i:02d}", price=10.0,
                       rating=4, in_stock=True, scraped_at="2026-03-01T00:00:00+00:00")
        db.upsert_book(conn, url=url, title=f"Item {i:02d}", price=9.0,
                       rating=4, in_stock=True, scraped_at="2026-03-02T00:00:00+00:00")
    wb = load_workbook(export_all(conn, tmp_path / "exports")["workbook"])
    conn.close()

    ws = wb["Price Trends"]
    assert ws.max_column == 11  # Date + TREND_MAX_PRODUCTS
    assert len(ws._charts[0].series) == 10


def test_history_csv_has_all_price_points(tmp_path, seeded_conn):
    paths = export_all(seeded_conn, tmp_path / "exports")
    with open(paths["history_csv"], newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 3  # Alpha x2, Beta x1


def test_print_report_summarises(seeded_conn, capsys):
    print_report(seeded_conn, limit=5)
    out = capsys.readouterr().out
    assert "Books tracked:      2" in out
    assert "Price changes seen: 1" in out
    assert "-20.0%" in out
